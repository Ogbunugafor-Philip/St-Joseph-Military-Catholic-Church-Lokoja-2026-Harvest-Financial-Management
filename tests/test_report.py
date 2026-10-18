import io
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from report import build_excel_report


class ReportTests(unittest.TestCase):
    def test_report_builds_from_synthetic_financial_data(self):
        with tempfile.TemporaryDirectory() as directory:
            db_path = Path(directory) / "synthetic.db"
            conn = sqlite3.connect(db_path)
            conn.executescript(
                """
                CREATE TABLE harvest_records (
                    id INTEGER PRIMARY KEY, name TEXT, phone_number TEXT,
                    harvest_activity TEXT, seed_sowing_category TEXT,
                    payment_mode TEXT, amount REAL, pledge_amount_paid REAL,
                    created_date TEXT
                );
                CREATE TABLE compulsory_contributions (
                    id INTEGER PRIMARY KEY, name TEXT, phone_number TEXT,
                    category TEXT, amount REAL, created_date TEXT
                );
                CREATE TABLE expenses (
                    id INTEGER PRIMARY KEY, expense_date TEXT,
                    description TEXT, amount REAL, approved_by TEXT
                );
                INSERT INTO harvest_records VALUES
                    (1, 'Synthetic Person', '0000000000', 'Harvest Proper', '', 'Pledge', 1000, 400, '2026-01-01'),
                    (2, 'Synthetic Person 2', '0000000001', 'Harvest Proper', '', 'Cash', 500, 500, '2026-01-02');
                INSERT INTO compulsory_contributions VALUES
                    (1, 'Synthetic Person', '0000000000', 'Man', 100, '2026-01-03');
                INSERT INTO expenses VALUES
                    (1, '2026-01-04', 'Synthetic expense', 50, 'Synthetic approver');
                """
            )
            conn.commit()
            conn.close()

            payload = build_excel_report(str(db_path))
            workbook = load_workbook(io.BytesIO(payload), data_only=True)
            self.assertIn("Summary", workbook.sheetnames)
            self.assertIn("All Harvest (Comprehensive)", workbook.sheetnames)
            self.assertIn("Harvest Proper", workbook.sheetnames)
            self.assertIn("Compulsory Levies", workbook.sheetnames)
            self.assertIn("Expenses", workbook.sheetnames)

            summary = workbook["Summary"]
            values = {
                summary.cell(row=row, column=1).value: summary.cell(row=row, column=4).value
                for row in range(1, summary.max_row + 1)
            }
            self.assertEqual(values["Total Harvest Pledged / Given (gross)"], 1500)
            self.assertEqual(values["Total Harvest Received (cash in hand)"], 900)
            self.assertEqual(values["Outstanding Pledges (yet to be redeemed)"], 600)
            self.assertEqual(values["Compulsory Levies Received"], 100)
            self.assertEqual(values["TOTAL INCOME RECEIVED"], 1000)
            self.assertEqual(values["TOTAL EXPENSES"], 50)
            self.assertEqual(values["NET BALANCE (income − expenses)"], 950)

    def test_xlsx_escapes_user_supplied_formula_text(self):
        with tempfile.TemporaryDirectory() as directory:
            db_path = Path(directory) / "synthetic.db"
            conn = sqlite3.connect(db_path)
            conn.executescript(
                """
                CREATE TABLE harvest_records (
                    id INTEGER PRIMARY KEY, name TEXT, phone_number TEXT,
                    harvest_activity TEXT, seed_sowing_category TEXT,
                    payment_mode TEXT, amount REAL, pledge_amount_paid REAL,
                    created_date TEXT
                );
                CREATE TABLE compulsory_contributions (
                    id INTEGER PRIMARY KEY, name TEXT, phone_number TEXT,
                    category TEXT, amount REAL, created_date TEXT
                );
                CREATE TABLE expenses (
                    id INTEGER PRIMARY KEY, expense_date TEXT,
                    description TEXT, amount REAL, approved_by TEXT
                );
                INSERT INTO harvest_records VALUES
                    (1, '=1+1', '@phone', 'Harvest Proper', '', 'Cash', 500, 500, '2026-01-02');
                """
            )
            conn.commit()
            conn.close()

            workbook = load_workbook(io.BytesIO(build_excel_report(str(db_path))), data_only=False)
            comprehensive = workbook["All Harvest (Comprehensive)"]
            self.assertEqual(comprehensive["B5"].value, "'=1+1")
            self.assertEqual(comprehensive["C5"].value, "'@phone")


if __name__ == "__main__":
    unittest.main()
