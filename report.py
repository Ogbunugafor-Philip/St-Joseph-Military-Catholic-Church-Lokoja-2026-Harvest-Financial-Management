# report.py — Professional multi-sheet Excel "Account Statement" builder.
# Pure / testable: no Streamlit imports. app.py calls build_excel_report(DB_PATH).

import io
import re
import sqlite3
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from export_safety import spreadsheet_safe_text

# ---- brand palette -------------------------------------------------------
NAVY = "1B2A4A"
NAVY_DK = "0F1A30"
GOLD = "B8860B"
LIGHT = "EEF1F7"
BAND = "F6F7FB"
GREEN = "1E7A3C"
RED = "A32222"
WHITE = "FFFFFF"
GREYTXT = "555555"

MONEY = '#,##0;[Red]-#,##0'
MONEY_NGN = '"₦" #,##0;[Red]-"₦" #,##0'  # ₦

CHURCH = "St Joseph Military Catholic Church, Lokoja"
TITLE = "2026 Harvest Account Statement"

thin = Side(style="thin", color="D0D5DD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

HARVEST_ACTIVITIES = [
    "Theme Selection", "Unveiling of Banner", "Seed Sowing", "Jesus’ Family",
    "Mary’s Family", "Joseph’s Family", "King and Queen Selection",
    "Children and Youths Harvest", "Societal Seed Sowing", "CMO/CWO Harvest",
    "Geopolitical Zones Harvest", "Sponsorship of Brochure Pictures",
    "Brochure Launch", "Harvest Proper",
]


# ---- styling helpers -----------------------------------------------------
def _fill(hexv):
    return PatternFill("solid", fgColor=hexv)


def _title_block(ws, ncols, subtitle):
    """Two merged banner rows at top; returns next free row (4)."""
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    c.value = CHURCH
    c.font = Font(name="Calibri", size=15, bold=True, color=WHITE)
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]
    c.value = subtitle
    c.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    c.fill = _fill(GOLD)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    return 4


def _header_row(ws, row, headers):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = _fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 26


def _autosize(ws, widths):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _data_table(ws, start_row, headers, rows, money_cols, widths, total_label_col=0):
    """Write a styled table. money_cols = set of 1-based column indexes to format as money.
    Returns (header_row, first_data_row, last_data_row)."""
    _header_row(ws, start_row, headers)
    r = start_row + 1
    first = r
    for i, rowvals in enumerate(rows):
        band = BAND if i % 2 else WHITE
        for j, v in enumerate(rowvals, start=1):
            c = ws.cell(row=r, column=j, value=spreadsheet_safe_text(v))
            c.border = BORDER
            c.fill = _fill(band)
            c.font = Font(size=9.5)
            if j in money_cols:
                c.number_format = MONEY_NGN
                c.alignment = Alignment(horizontal="right")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        r += 1
    last = r - 1
    _autosize(ws, widths)
    ws.freeze_panes = ws.cell(row=first, column=1)
    return start_row, first, last


def _total_row(ws, row, ncols, label, label_col, totals, money_cols):
    """totals: dict {col_index: value}."""
    for j in range(1, ncols + 1):
        c = ws.cell(row=row, column=j)
        c.fill = _fill(NAVY_DK)
        c.border = BORDER
        c.font = Font(bold=True, color=WHITE, size=10)
        if j == label_col:
            c.value = label
            c.alignment = Alignment(horizontal="left")
        if j in totals:
            c.value = totals[j]
            if j in money_cols:
                c.number_format = MONEY_NGN
            c.alignment = Alignment(horizontal="right")
    ws.row_dimensions[row].height = 20


def _safe_sheet_name(name, used):
    s = re.sub(r"[:\\/?*\[\]]", " ", name).strip()
    s = s[:31]
    base = s
    n = 2
    while s.lower() in used:
        suffix = f" ({n})"
        s = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(s.lower())
    return s


# ---- data loading --------------------------------------------------------
def _load(db_path):
    conn = sqlite3.connect(db_path)
    h = pd.read_sql_query("SELECT * FROM harvest_records ORDER BY id", conn)
    c = pd.read_sql_query("SELECT * FROM compulsory_contributions ORDER BY id", conn)
    e = pd.read_sql_query("SELECT * FROM expenses ORDER BY id", conn)
    conn.close()

    if not h.empty:
        h["balance"] = h["amount"] - h["pledge_amount_paid"]
        h["status"] = h.apply(
            lambda r: "Paid" if (r["payment_mode"] in ("Cash", "Transfer") or r["balance"] <= 0) else "Unpaid",
            axis=1,
        )
    else:
        h["balance"] = []
        h["status"] = []
    return h, c, e


# ---- summary sheet -------------------------------------------------------
def _summary_sheet(ws, h, c, e, today):
    H_amt = float(h["amount"].sum()) if not h.empty else 0.0
    H_paid = float(h["pledge_amount_paid"].sum()) if not h.empty else 0.0
    H_bal = H_amt - H_paid
    C_amt = float(c["amount"].sum()) if not c.empty else 0.0
    E_amt = float(e["amount"].sum()) if not e.empty else 0.0
    income = H_paid + C_amt
    net = income - E_amt

    ws.sheet_view.showGridLines = False
    row = _title_block(ws, 5, f"{TITLE}  —  Executive Summary")
    ws.cell(row=row, column=1, value=f"Statement generated: {today.strftime('%d %B %Y')}").font = Font(
        italic=True, size=9, color=GREYTXT)
    row += 2

    # ---- Financial position block
    ws.cell(row=row, column=1, value="FINANCIAL POSITION").font = Font(bold=True, size=12, color=NAVY)
    row += 1
    fin = [
        ("Total Harvest Pledged / Given (gross)", H_amt, NAVY),
        ("Total Harvest Received (cash in hand)", H_paid, GREEN),
        ("Outstanding Pledges (yet to be redeemed)", H_bal, RED),
        ("Compulsory Levies Received", C_amt, GREEN),
        ("TOTAL INCOME RECEIVED", income, GREEN),
        ("TOTAL EXPENSES", E_amt, RED),
        ("NET BALANCE (income − expenses)", net, GREEN if net >= 0 else RED),
    ]
    for i, (label, val, col) in enumerate(fin):
        emphatic = label.isupper()
        lc = ws.cell(row=row, column=1, value=label)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        vc = ws.cell(row=row, column=4, value=val)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        band = LIGHT if not emphatic else NAVY
        lc.fill = _fill(band)
        vc.fill = _fill(band)
        lc.border = BORDER
        vc.border = BORDER
        vc.number_format = MONEY_NGN
        vc.alignment = Alignment(horizontal="right", vertical="center")
        lc.alignment = Alignment(horizontal="left", vertical="center")
        if emphatic:
            lc.font = Font(bold=True, color=WHITE, size=11)
            vc.font = Font(bold=True, color=WHITE, size=11)
        else:
            lc.font = Font(size=10, color="333333")
            vc.font = Font(bold=True, size=10, color=col)
        ws.row_dimensions[row].height = 19
        row += 1
    row += 1

    # ---- By activity table
    act_start = row
    ws.cell(row=row, column=1, value="HARVEST BY PROGRAMME / ACTIVITY").font = Font(bold=True, size=12, color=NAVY)
    row += 1
    headers = ["Activity / Programme", "Records", "Amount", "Received", "Balance"]
    if not h.empty:
        g = h.groupby("harvest_activity").agg(
            records=("id", "count"), amount=("amount", "sum"),
            received=("pledge_amount_paid", "sum")).reset_index()
        g["balance"] = g["amount"] - g["received"]
        g = g.sort_values("amount", ascending=False)
        rows = [[r["harvest_activity"], int(r["records"]), float(r["amount"]),
                 float(r["received"]), float(r["balance"])] for _, r in g.iterrows()]
    else:
        rows = []
    hr, first, last = _data_table(
        ws, row, headers, rows, money_cols={3, 4, 5},
        widths=[34, 10, 16, 16, 16])
    ws.freeze_panes = "A4"  # keep title visible, not mid-sheet
    tot_row = last + 1
    _total_row(ws, tot_row, 5, "TOTAL", 1,
               {2: int(sum(r[1] for r in rows)) if rows else 0,
                3: sum(r[2] for r in rows), 4: sum(r[3] for r in rows),
                5: sum(r[4] for r in rows)}, money_cols={3, 4, 5})

    # bar chart: Amount vs Received by activity
    if rows:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Amount vs Received by Activity (₦)"
        chart.height = 0.55 * (last - first + 2) + 3
        chart.width = 18
        data = Reference(ws, min_col=3, max_col=4, min_row=hr, max_row=last)
        cats = Reference(ws, min_col=1, max_col=1, min_row=first, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.gapWidth = 60
        ws.add_chart(chart, f"A{tot_row + 2}")

    # ---- By payment mode + compulsory category (placed to the right, col 7)
    if not h.empty:
        pm = h.groupby("payment_mode").agg(
            records=("id", "count"), amount=("amount", "sum"),
            received=("pledge_amount_paid", "sum")).reset_index()
        prow = act_start
        ws.cell(row=prow, column=7, value="BY PAYMENT MODE").font = Font(bold=True, size=12, color=NAVY)
        prow += 1
        ph = ["Mode", "Records", "Amount", "Received"]
        prows = [[r["payment_mode"], int(r["records"]), float(r["amount"]), float(r["received"])]
                 for _, r in pm.iterrows()]
        for j, htxt in enumerate(ph, start=7):
            cc = ws.cell(row=prow, column=j, value=htxt)
            cc.font = Font(bold=True, color=WHITE, size=10)
            cc.fill = _fill(NAVY)
            cc.alignment = Alignment(horizontal="center", vertical="center")
            cc.border = BORDER
        rr = prow + 1
        for i, pv in enumerate(prows):
            band = BAND if i % 2 else WHITE
            for j, v in enumerate(pv, start=7):
                cc = ws.cell(row=rr, column=j, value=v)
                cc.border = BORDER
                cc.fill = _fill(band)
                cc.font = Font(size=9.5)
                if j in (9, 10):
                    cc.number_format = MONEY_NGN
                    cc.alignment = Alignment(horizontal="right")
            rr += 1
        for j, w in zip(range(7, 11), [12, 10, 15, 15]):
            ws.column_dimensions[get_column_letter(j)].width = w

        # pie chart of received by mode
        pie = PieChart()
        pie.title = "Received by Payment Mode"
        pdata = Reference(ws, min_col=10, min_row=prow, max_row=rr - 1)
        plabels = Reference(ws, min_col=7, min_row=prow + 1, max_row=rr - 1)
        pie.add_data(pdata, titles_from_data=True)
        pie.set_categories(plabels)
        pie.height = 7
        pie.width = 11
        ws.add_chart(pie, f"G{rr + 1}")

    # category summary far right
    if not c.empty:
        cg = c.groupby("category").agg(records=("id", "count"), amount=("amount", "sum")).reset_index()
        crow = act_start
        ws.cell(row=crow, column=12, value="COMPULSORY LEVIES").font = Font(bold=True, size=12, color=NAVY)
        crow += 1
        for j, htxt in enumerate(["Category", "Records", "Amount"], start=12):
            cc = ws.cell(row=crow, column=j, value=htxt)
            cc.font = Font(bold=True, color=WHITE, size=10)
            cc.fill = _fill(NAVY)
            cc.alignment = Alignment(horizontal="center")
            cc.border = BORDER
        rr = crow + 1
        for i, (_, r) in enumerate(cg.iterrows()):
            band = BAND if i % 2 else WHITE
            vals = [r["category"], int(r["records"]), float(r["amount"])]
            for j, v in enumerate(vals, start=12):
                cc = ws.cell(row=rr, column=j, value=v)
                cc.border = BORDER
                cc.fill = _fill(band)
                cc.font = Font(size=9.5)
                if j == 14:
                    cc.number_format = MONEY_NGN
                    cc.alignment = Alignment(horizontal="right")
            rr += 1
        for j, w in zip(range(12, 15), [14, 10, 15]):
            ws.column_dimensions[get_column_letter(j)].width = w


# ---- detail sheet builders ----------------------------------------------
def _harvest_detail_rows(df):
    rows = []
    for _, r in df.iterrows():
        rows.append([
            r["created_date"], r["name"], str(r["phone_number"] or ""),
            r["seed_sowing_category"] or "", r["payment_mode"],
            float(r["amount"]), float(r["pledge_amount_paid"]),
            float(r["balance"]), r["status"],
        ])
    return rows


def _activity_sheet(ws, df, activity):
    ws.sheet_view.showGridLines = False
    _title_block(ws, 9, f"{TITLE}  —  {activity}")
    headers = ["Date", "Name", "Phone", "Seed Category", "Payment Mode",
               "Amount", "Amount Paid", "Balance", "Status"]
    rows = _harvest_detail_rows(df)
    hr, first, last = _data_table(
        ws, 4, headers, rows, money_cols={6, 7, 8},
        widths=[12, 26, 16, 14, 14, 14, 14, 14, 10])
    _total_row(ws, last + 1, 9, "TOTAL", 1,
               {6: sum(r[5] for r in rows), 7: sum(r[6] for r in rows),
                8: sum(r[7] for r in rows)}, money_cols={6, 7, 8})


def _comprehensive_sheet(ws, h, c, e):
    ws.sheet_view.showGridLines = False
    _title_block(ws, 10, f"{TITLE}  —  All Harvest Records (Comprehensive)")
    headers = ["Date", "Name", "Phone", "Activity / Programme", "Seed Category",
               "Payment Mode", "Amount", "Amount Paid", "Balance", "Status"]
    rows = []
    for _, r in h.iterrows():
        rows.append([
            r["created_date"], r["name"], str(r["phone_number"] or ""),
            r["harvest_activity"], r["seed_sowing_category"] or "",
            r["payment_mode"], float(r["amount"]), float(r["pledge_amount_paid"]),
            float(r["balance"]), r["status"],
        ])
    hr, first, last = _data_table(
        ws, 4, headers, rows, money_cols={7, 8, 9},
        widths=[12, 26, 16, 24, 13, 13, 14, 14, 14, 10])
    total_row = last + 1
    _total_row(ws, total_row, 10, "TOTAL", 1,
               {7: sum(r[6] for r in rows), 8: sum(r[7] for r in rows),
                9: sum(r[8] for r in rows)}, money_cols={7, 8, 9})

    # ---- financial summary block at the bottom of the sheet ----
    h_pledged = sum(r[6] for r in rows)              # total amount pledged / given
    h_redeemed = sum(r[7] for r in rows)             # total amount actually paid in
    h_outstanding = sum(r[8] for r in rows)          # pledged not yet redeemed
    c_amt = float(c["amount"].sum()) if not c.empty else 0.0
    e_amt = float(e["amount"].sum()) if not e.empty else 0.0
    total_income = h_redeemed + c_amt                # cash actually received
    balance_at_hand = total_income - e_amt

    srow = total_row + 2
    # heading bar (merged across the table width)
    ws.merge_cells(start_row=srow, start_column=1, end_row=srow, end_column=10)
    hd = ws.cell(row=srow, column=1, value="ACCOUNT SUMMARY")
    hd.font = Font(bold=True, color=WHITE, size=12)
    hd.fill = _fill(NAVY)
    hd.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[srow].height = 24
    srow += 1

    summary = [
        ("Total Harvest Pledged / Given (gross)", h_pledged, "333333", LIGHT, False),
        ("Total Pledges Redeemed (paid in)", h_redeemed, GREEN, LIGHT, False),
        ("Outstanding Pledges (not yet redeemed)", h_outstanding, RED, LIGHT, False),
        ("Compulsory Levies Received", c_amt, GREEN, LIGHT, False),
        ("TOTAL INCOME RECEIVED  (redeemed + levies)", total_income, WHITE, NAVY, True),
        ("TOTAL EXPENSES", e_amt, RED, LIGHT, False),
        ("BALANCE AT HAND  (income − expenses)", balance_at_hand,
         WHITE, (GREEN if balance_at_hand >= 0 else RED), True),
    ]
    for label, val, txt_col, fill_col, emphatic in summary:
        ws.merge_cells(start_row=srow, start_column=1, end_row=srow, end_column=7)
        lc = ws.cell(row=srow, column=1, value=label)
        ws.merge_cells(start_row=srow, start_column=8, end_row=srow, end_column=10)
        vc = ws.cell(row=srow, column=8, value=val)
        for cc in (lc, vc):
            cc.fill = _fill(fill_col)
            cc.border = BORDER
        lc.alignment = Alignment(horizontal="left", vertical="center")
        lc.font = Font(bold=emphatic, size=11 if emphatic else 10, color=txt_col)
        vc.number_format = MONEY_NGN
        vc.alignment = Alignment(horizontal="right", vertical="center")
        vc.font = Font(bold=True, size=11 if emphatic else 10, color=txt_col)
        ws.row_dimensions[srow].height = 20 if emphatic else 18
        srow += 1


def _compulsory_sheet(ws, c):
    ws.sheet_view.showGridLines = False
    _title_block(ws, 5, f"{TITLE}  —  Compulsory Levies")
    headers = ["Date", "Name", "Phone", "Category", "Amount"]
    rows = [[r["created_date"], r["name"], str(r["phone_number"] or ""),
             r["category"], float(r["amount"])] for _, r in c.iterrows()]
    _data_table(ws, 4, headers, rows, money_cols={5}, widths=[12, 28, 16, 14, 16])
    _total_row(ws, 4 + 1 + len(rows), 5, "TOTAL", 1,
               {5: sum(r[4] for r in rows)} if rows else {5: 0}, money_cols={5})


def _expenses_sheet(ws, e):
    ws.sheet_view.showGridLines = False
    _title_block(ws, 4, f"{TITLE}  —  Expenses")
    headers = ["Date", "Description", "Amount", "Approved By"]
    rows = [[r["expense_date"], r["description"], float(r["amount"]), r["approved_by"] or ""]
            for _, r in e.iterrows()]
    _data_table(ws, 4, headers, rows, money_cols={3}, widths=[14, 48, 16, 22])
    _total_row(ws, 4 + 1 + len(rows), 4, "TOTAL", 1,
               {3: sum(r[2] for r in rows)} if rows else {3: 0}, money_cols={3})


# ---- public entry point --------------------------------------------------
def build_excel_report(db_path, today=None):
    """Build the full workbook and return raw .xlsx bytes."""
    today = today or date.today()
    h, c, e = _load(db_path)

    wb = Workbook()
    used = set()

    ws = wb.active
    ws.title = _safe_sheet_name("Summary", used)
    _summary_sheet(ws, h, c, e, today)

    # comprehensive
    ws = wb.create_sheet(_safe_sheet_name("All Harvest (Comprehensive)", used))
    _comprehensive_sheet(ws, h, c, e)

    # one sheet per activity that has data (in canonical order)
    present = list(h["harvest_activity"].dropna().unique()) if not h.empty else []
    ordered = [a for a in HARVEST_ACTIVITIES if a in present] + \
              [a for a in present if a not in HARVEST_ACTIVITIES]
    for activity in ordered:
        sub = h[h["harvest_activity"] == activity]
        ws = wb.create_sheet(_safe_sheet_name(activity, used))
        _activity_sheet(ws, sub, activity)

    # compulsory + expenses
    ws = wb.create_sheet(_safe_sheet_name("Compulsory Levies", used))
    _compulsory_sheet(ws, c)
    ws = wb.create_sheet(_safe_sheet_name("Expenses", used))
    _expenses_sheet(ws, e)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


REPORT_FILENAME = "St Joseph Military Catholic Church 2026 Harvest Account Statement.xlsx"


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Build an Excel statement from an explicitly selected database.")
    parser.add_argument("database", help="Path to the SQLite database to read")
    parser.add_argument("--output", default=REPORT_FILENAME, help="Output workbook path")
    args = parser.parse_args()

    data = build_excel_report(args.database)
    with open(args.output, "wb") as f:
        f.write(data)
    print("WROTE", args.output, len(data), "bytes")
